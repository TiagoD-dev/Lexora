from io import BytesIO
import re
import os
from pathlib import Path
from uuid import uuid4

from alembic import command
from alembic.config import Config
from fastapi import Depends, FastAPI, File, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse

from sqlalchemy.orm import Session

from . import models
from .db import get_db
from .routers import auth, billing, cases, clients, legal_updates, notifications, portal
from .security import get_current_user, get_user_from_token

API_ROOT = Path(__file__).resolve().parent.parent
alembic_cfg = Config(str(API_ROOT / "alembic.ini"))
alembic_cfg.set_main_option("script_location", str(API_ROOT / "alembic"))
command.upgrade(alembic_cfg, "head")

app = FastAPI(
    title="Lexora API",
    description="API base do assistente jurídico digital.",
    version="0.1.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[origin.strip() for origin in os.environ.get("LEXORA_CORS_ORIGINS", "http://localhost:8081").split(",") if origin.strip()],
    allow_origin_regex=r"^http://(localhost|127\.0\.0\.1)(:\d+)?$",
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.middleware("http")
async def private_portal_responses(request, call_next):
    response = await call_next(request)
    if request.url.path.startswith("/portal/"):
        response.headers["Cache-Control"] = "no-store"
        response.headers["X-Content-Type-Options"] = "nosniff"
    return response


app.include_router(portal.router)
app.include_router(auth.router)
app.include_router(billing.router)
app.include_router(clients.router)
app.include_router(cases.router)
app.include_router(legal_updates.router)
app.include_router(notifications.router)


@app.get("/health", tags=["system"])
def health_check() -> dict[str, str]:
    return {"status": "ok", "service": "lexora-api"}


MAX_FILE_SIZE = 25 * 1024 * 1024
DOCUMENTS_DIR = Path(__file__).resolve().parent.parent / "documents_storage"
DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)


def extract_text(filename: str, content: bytes) -> tuple[str, int | None]:
    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(content))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages), len(reader.pages)
    if extension in {".docx", ".doc"}:
        if extension == ".doc":
            raise HTTPException(422, "O formato DOC antigo deve ser convertido para DOCX.")
        from docx import Document
        document = Document(BytesIO(content))
        return "\n".join(paragraph.text for paragraph in document.paragraphs), None
    if extension in {".xlsx", ".xls"}:
        if extension == ".xls":
            raise HTTPException(422, "O formato XLS antigo deve ser convertido para XLSX.")
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        rows = []
        for sheet in workbook.worksheets:
            rows.append(f"[{sheet.title}]")
            rows.extend(" | ".join(str(value) for value in row if value is not None) for row in sheet.iter_rows(values_only=True))
        return "\n".join(rows), None
    if extension in {".txt", ".md", ".csv"}:
        return content.decode("utf-8", errors="replace"), None
    raise HTTPException(415, "Formato ainda não suportado para extração de texto.")


def build_suggestions(text: str) -> list[dict]:
    clean = re.sub(r"[ \t]+", " ", text).strip()
    suggestions: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()

    def add(kind: str, value: str, detail: str = "", start: int | None = None, end: int | None = None) -> None:
        normalized = value.strip(" \n\t,;:.-")
        key = (kind, normalized.casefold())
        if len(normalized) < 3 or key in seen:
            return
        seen.add(key)
        excerpt = ""
        if start is not None and end is not None:
            excerpt_start = max(0, start - 120)
            excerpt_end = min(len(clean), end + 120)
            excerpt = clean[excerpt_start:excerpt_end].strip()
        suggestions.append({
            "id": str(uuid4()), "type": kind, "value": normalized[:500], "detail": detail,
            "excerpt": excerpt[:800], "characterStart": start, "characterEnd": end,
        })

    for match in re.finditer(r"\b(?:0?[1-9]|[12]\d|3[01])[/-](?:0?[1-9]|1[0-2])[/-](?:19|20)\d{2}\b", clean):
        add("Data", match.group(), "Data encontrada no documento", match.start(), match.end())
    for match in re.finditer(r"\b[A-ZÁÀÂÃÉÊÍÓÔÕÚÇ][a-záàâãéêíóôõúç]+(?:\s+(?:da|de|do|dos|das|e)?\s*[A-ZÁÀÂÃÉÊÍÓÔÕÚÇ][a-záàâãéêíóôõúç]+){1,3}\b", clean):
        add("Entidade", match.group(), "Pessoa ou entidade potencial; confirmar papel", match.start(), match.end())
        if sum(item["type"] == "Entidade" for item in suggestions) >= 8:
            break
    sentences = [sentence.strip() for sentence in re.split(r"(?<=[.!?])\s+|\n+", clean) if 35 <= len(sentence.strip()) <= 500]
    for sentence in sentences[:10]:
        start = clean.find(sentence)
        add("Facto", sentence, "Excerto proposto como facto; confirmar conteúdo e estado", start if start >= 0 else None, start + len(sentence) if start >= 0 else None)
    return suggestions[:24]


@app.post("/documents/extract", tags=["documents"])
async def extract_document(
    file: UploadFile = File(...),
    current_user: models.User = Depends(get_current_user),
) -> dict:
    content = await file.read(MAX_FILE_SIZE + 1)
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(413, "O ficheiro excede o limite de 25 MB.")
    if not content:
        raise HTTPException(400, "O ficheiro está vazio.")
    try:
        text, pages = extract_text(file.filename or "documento", content)
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(422, f"Não foi possível extrair o documento: {error}") from error
    if not text.strip():
        raise HTTPException(422, "Não foi encontrado texto. O documento pode necessitar de OCR.")
    extension = Path(file.filename or "").suffix.lower()[:10]
    stored_name = f"{current_user.id}_{uuid4().hex}{extension}"
    (DOCUMENTS_DIR / stored_name).write_bytes(content)
    (DOCUMENTS_DIR / f"{stored_name}.txt").write_text(text, encoding="utf-8")  # usado pelo assistente
    return {
        "text": text[:250_000],
        "characterCount": len(text),
        "pageCount": pages,
        "suggestions": build_suggestions(text),
        "fileId": stored_name,
    }


@app.get("/documents/file/{stored_name}", tags=["documents"])
def get_document_file(stored_name: str, token: str = Query(...), db: Session = Depends(get_db)) -> FileResponse:
    if "/" in stored_name or "\\" in stored_name or ".." in stored_name:
        raise HTTPException(404, "Ficheiro não encontrado.")
    user = get_user_from_token(token, db)
    if not stored_name.startswith(f"{user.id}_"):
        raise HTTPException(404, "Ficheiro não encontrado.")
    path = DOCUMENTS_DIR / stored_name
    if not path.is_file():
        raise HTTPException(404, "Ficheiro não encontrado.")
    return FileResponse(path)
